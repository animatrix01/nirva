"""Export endpoints (E1/E2): PPTX / Word / Excel investigation reports.

Files are built on demand from live report data (same sources as the UI)
into the runtime reports dir, then streamed with FileResponse. Visibility
reuses the jobs _can_view check; per-chunk allowed_roles filtering happens
inside the data collection, never in the client.
"""

from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from app.db import get_db
from app.deps import CurrentUser
from app.exporters.report_data import collect, reports_dir
from app.models import Document, Job
from app.routers.jobs import _can_view

router = APIRouter(prefix="/api/jobs", tags=["exports"])


def _job_or_403(db: Session, user: CurrentUser, job_id: str) -> Job:
    job = db.get(Job, job_id)
    if job is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Job not found")
    if not _can_view(user, job):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden")
    return job


@router.get("/{job_id}/export/pptx")
def export_pptx(
    job_id: str,
    user: CurrentUser,
    db: Annotated[Session, Depends(get_db)],
) -> FileResponse:
    from app.exporters.pptx import build_pptx

    job = _job_or_403(db, user, job_id)
    data = collect(db, job, user.role)
    out_path = f"{reports_dir()}/{job.id}.pptx"
    build_pptx(data, out_path)
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename=f"{data.equipment_id}-investigation-{job.id[:8]}.pptx",
    )


@router.get("/{job_id}/export/word")
def export_word(
    job_id: str,
    user: CurrentUser,
    db: Annotated[Session, Depends(get_db)],
) -> FileResponse:
    from app.exporters.docx import build_docx

    job = _job_or_403(db, user, job_id)
    data = collect(db, job, user.role)
    out_path = f"{reports_dir()}/{job.id}.docx"
    build_docx(data, out_path)
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"{data.equipment_id}-investigation-{job.id[:8]}.docx",
    )


@router.get("/{job_id}/export/excel")
def export_excel(
    job_id: str,
    user: CurrentUser,
    db: Annotated[Session, Depends(get_db)],
) -> FileResponse:
    from app.exporters.xlsx import build_xlsx

    job = _job_or_403(db, user, job_id)
    data = collect(db, job, user.role)
    out_path = f"{reports_dir()}/{job.id}.xlsx"
    build_xlsx(data, out_path, code=data.working_code, result=data.working_code_result)
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"{data.equipment_id}-investigation-{job.id[:8]}.xlsx",
    )


documents_router = APIRouter(prefix="/api/documents", tags=["exports"])
"""Second router for the library-level export (registered alongside)."""


@documents_router.get("/export/excel")
def export_documents_excel(
    user: CurrentUser,
    db: Annotated[Session, Depends(get_db)],
    department: str | None = None,
    document_type: str | None = None,
) -> FileResponse:
    """Export the visible (ACL-filtered) document list, optionally narrowed
    by department/document_type."""
    import openpyxl  # type: ignore[import-untyped]

    docs = db.query(Document).order_by(Document.uploaded_at.desc()).all()
    rows = []
    for doc in docs:
        if user.role not in (doc.allowed_roles or []):
            continue
        if department and doc.department != department:
            continue
        if document_type and doc.document_type != document_type:
            continue
        rows.append(doc)

    out_path = f"{reports_dir()}/documents-{user.id[:8]}.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Documents"
    headers = ["Title", "Type", "Revision", "Department", "Equipment", "Status", "Uploaded"]
    for col, name in enumerate(headers, start=1):
        sheet.cell(row=1, column=col, value=name)
    for idx, doc in enumerate(rows, start=1):
        sheet.cell(row=idx + 1, column=1, value=doc.title)
        sheet.cell(row=idx + 1, column=2, value=doc.document_type)
        sheet.cell(row=idx + 1, column=3, value=doc.revision)
        sheet.cell(row=idx + 1, column=4, value=doc.department)
        sheet.cell(row=idx + 1, column=5, value=doc.equipment_id or "")
        sheet.cell(row=idx + 1, column=6, value=doc.status)
        uploaded = doc.uploaded_at
        sheet.cell(row=idx + 1, column=7, value=uploaded.isoformat() if uploaded else "")
    sheet.column_dimensions["A"].width = 42
    wb.save(out_path)
    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="nirva-documents.xlsx",
    )
