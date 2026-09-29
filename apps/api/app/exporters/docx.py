"""DOCX exporter (E2): investigation report document, same sections as PPT.

Cover / Primary Finding / Supporting Evidence / Contradiction /
Recommended Action / Evidence Appendix.
"""

from __future__ import annotations

from datetime import date

from app.exporters.report_data import ReportData
from docx import Document as _new_document
from docx.document import Document
from docx.shared import Pt, RGBColor

SLATE = RGBColor(0x0F, 0x17, 0x2A)


def _title(doc: Document, text: str) -> None:
    paragraph = doc.add_heading(text, level=0)
    for run in paragraph.runs:
        run.font.color.rgb = SLATE
        run.font.bold = True


def _h1(doc: Document, text: str) -> None:
    paragraph = doc.add_heading(text, level=1)
    for run in paragraph.runs:
        run.font.color.rgb = SLATE
        run.font.bold = True


def _body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    for run in paragraph.runs:
        run.font.size = Pt(11)


def _bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text, style="List Bullet")
    for run in paragraph.runs:
        run.font.size = Pt(11)


def build_docx(data: ReportData, out_path: str) -> str:
    doc: Document = _new_document()

    _title(doc, "Investigation Report")
    _body(doc, "ForgeOS — Team CodeWizards")
    _body(doc, f"Equipment {data.equipment_id}")
    _body(doc, data.objective)
    when = data.created_at.date().isoformat() if data.created_at else date.today().isoformat()
    _body(doc, f"Job {data.job_id[:8]} · {when} · status: {data.status}")

    _h1(doc, "Primary Finding")
    _body(doc, data.findings or "No findings recorded.")
    _body(doc, f"Confidence: {data.confidence:.0%} · {len(data.evidence)} sources")

    _h1(doc, "Supporting Evidence")
    if not data.evidence:
        _body(doc, "No indexed evidence was available.")
    for idx, item in enumerate(data.evidence[:8], start=1):
        _body(
            doc,
            f"{idx}. {item.title} — {item.locator} (score {item.score:.0%})\n"
            f"{item.passage[:400]}",
        )

    _h1(doc, "Contradictions")
    if not data.contradictions:
        _body(doc, "No cross-source conflicts detected.")
    for row in data.contradictions:
        _body(doc, f"Topic: {row.topic}")
        _body(doc, f"A — {row.source_a}: {row.claim_a[:400]}")
        _body(doc, f"B — {row.source_b}: {row.claim_b[:400]}")

    _h1(doc, "Recommended Action")
    if data.recommendation:
        _body(doc, data.recommendation[:1200])
        _body(doc, f"Action: {data.recommendation_action} · Risk: {data.recommendation_risk or '—'}")
    else:
        _body(doc, "No work-order proposal was created for this investigation.")
    _body(doc, "Consequential actions require supervisor approval before execution.")

    _h1(doc, "Evidence Appendix")
    if not data.citations:
        _body(doc, "No citations recorded.")
    for idx, citation in enumerate(data.citations, start=1):
        _body(
            doc,
            f"[{idx}] {citation.get('title', '')} · Rev {citation.get('revision', '')} "
            f"· {citation.get('locator', '')} · {citation.get('chunk_id', '')}",
        )
    if data.assumptions:
        _h1(doc, "Assumptions (labeled, not evidenced)")
        for assumption in data.assumptions:
            _bullet(doc, str(assumption)[:200])
    if data.gaps:
        _h1(doc, "Gaps")
        for gap in data.gaps:
            _bullet(doc, str(gap)[:200])

    doc.save(out_path)
    return out_path
