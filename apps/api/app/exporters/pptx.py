"""PPTX exporter (E1): investigation report deck from live report data.

Theme per team spec: white background, slate-900 (#0F172A) text, bold
titles, lime (#C8F135) accents. Six slides: Cover / Primary Finding /
Supporting 1-4 / Contradiction / Recommended / Evidence Appendix.
"""

from __future__ import annotations

from datetime import date

from app.exporters.report_data import ContradictionRow, EvidenceRow, ReportData
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.util import Inches, Pt

SLATE = RGBColor(0x0F, 0x17, 0x2A)
LIME = RGBColor(0xC8, 0xF1, 0x35)
GRAY = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE_FILL = RGBColor(0xFF, 0xF7, 0xED)
ORANGE_BORDER = RGBColor(0xFB, 0x92, 0x3C)
EMERALD_FILL = RGBColor(0xEC, 0xFD, 0xF5)
EMERALD_BORDER = RGBColor(0x10, 0xB9, 0x81)

TITLE_FONT_SIZE = Pt(28)
BODY_FONT_SIZE = Pt(14)
SMALL_FONT_SIZE = Pt(11)


def _white_bg(slide: Slide) -> None:
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE


def _textbox(
    slide: Slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    size: Pt = BODY_FONT_SIZE,
    bold: bool = False,
    color: RGBColor = SLATE,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = alignment
    run = paragraph.runs[0]
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color


def _accent_bar(slide: Slide, top: float = 0.0, height: float = 0.12) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(top), Inches(10.0), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIME
    shape.line.fill.background()


def _new_slide(prs: PresentationType) -> Slide:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _white_bg(slide)
    _accent_bar(slide)
    return slide


def _slide_cover(prs: PresentationType, data: ReportData) -> None:
    slide = _new_slide(prs)
    _textbox(slide, 0.7, 1.2, 8.6, 0.5, "ForgeOS — Team CodeWizards",
             size=Pt(16), bold=True, color=GRAY)
    _textbox(slide, 0.7, 1.8, 8.6, 1.0, "Investigation Report",
             size=Pt(40), bold=True, color=SLATE)
    _textbox(slide, 0.7, 3.0, 8.6, 0.6, f"Equipment {data.equipment_id}",
             size=TITLE_FONT_SIZE, bold=True, color=SLATE)
    objective = data.objective[:220] + ("…" if len(data.objective) > 220 else "")
    _textbox(slide, 0.7, 3.9, 8.6, 1.6, objective, size=BODY_FONT_SIZE, color=GRAY)
    when = data.created_at.date().isoformat() if data.created_at else date.today().isoformat()
    _textbox(slide, 0.7, 5.8, 8.6, 0.5,
             f"Job {data.job_id[:8]} · {when} · status: {data.status}",
             size=SMALL_FONT_SIZE, color=GRAY)


def _slide_finding(prs: PresentationType, data: ReportData) -> None:
    slide = _new_slide(prs)
    _textbox(slide, 0.7, 0.5, 8.6, 0.7, "Primary Finding",
             size=TITLE_FONT_SIZE, bold=True, color=SLATE)
    findings = data.findings[:900] + ("…" if len(data.findings) > 900 else "")
    _textbox(slide, 0.7, 1.5, 8.6, 3.6, findings or "No findings recorded.",
             size=BODY_FONT_SIZE, color=SLATE)
    _textbox(slide, 0.7, 5.4, 8.6, 0.6,
             f"Confidence: {data.confidence:.0%}  ·  {len(data.evidence)} sources",
             size=BODY_FONT_SIZE, bold=True, color=SLATE)


def _slide_supporting(prs: PresentationType, data: ReportData) -> None:
    slide = _new_slide(prs)
    _textbox(slide, 0.7, 0.5, 8.6, 0.7, "Supporting Evidence 1–4",
             size=TITLE_FONT_SIZE, bold=True, color=SLATE)
    items: list[EvidenceRow] = data.evidence[:4]
    if not items:
        _textbox(slide, 0.7, 1.6, 8.6, 1.0, "No indexed evidence was available.",
                 size=BODY_FONT_SIZE, color=GRAY)
        return
    top = 1.6
    for idx, item in enumerate(items, start=1):
        head = f"{idx}. {item.title} — {item.locator}  (score {item.score:.0%})"
        _textbox(slide, 0.7, top, 8.6, 0.5, head, size=Pt(13), bold=True, color=SLATE)
        passage = item.passage[:260] + ("…" if len(item.passage) > 260 else "")
        _textbox(slide, 0.7, top + 0.5, 8.6, 0.7, passage, size=SMALL_FONT_SIZE, color=GRAY)
        top += 1.25


def _contradiction_table(slide: Slide, row: ContradictionRow, top: float) -> None:
    table_shape = slide.shapes.add_table(3, 2, Inches(0.7), Inches(top), Inches(8.6), Inches(2.4))
    table = table_shape.table
    headers = ("Source A", "Source B")
    bodies = (
        f"{row.source_a}\n{row.claim_a[:220]}",
        f"{row.source_b}\n{row.claim_b[:220]}",
    )
    for col in range(2):
        cell = table.cell(0, col)
        cell.text = headers[col]
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(13)
                run.font.color.rgb = SLATE
        cell = table.cell(1, col)
        cell.text = bodies[col]
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = SMALL_FONT_SIZE
                run.font.color.rgb = SLATE
        table.cell(2, col).text = ""
    for r in range(3):
        for c in range(2):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ORANGE_FILL


def _slide_contradiction(prs: PresentationType, data: ReportData) -> None:
    slide = _new_slide(prs)
    _textbox(slide, 0.7, 0.5, 8.6, 0.7, "Contradiction",
             size=TITLE_FONT_SIZE, bold=True, color=SLATE)
    if not data.contradictions:
        _textbox(slide, 0.7, 1.6, 8.6, 1.0,
                 "No cross-source conflicts detected for this investigation.",
                 size=BODY_FONT_SIZE, color=GRAY)
        return
    row = data.contradictions[0]
    _textbox(slide, 0.7, 1.3, 8.6, 0.5, row.topic, size=Pt(16), bold=True, color=SLATE)
    _contradiction_table(slide, row, top=1.9)
    _textbox(slide, 0.7, 5.9, 8.6, 0.5,
             "Conflict surfaced for human verification — not silently resolved.",
             size=SMALL_FONT_SIZE, color=GRAY)


def _slide_recommended(prs: PresentationType, data: ReportData) -> None:
    slide = _new_slide(prs)
    _textbox(slide, 0.7, 0.5, 8.6, 0.7, "Recommended Action",
             size=TITLE_FONT_SIZE, bold=True, color=SLATE)
    text = data.recommendation[:800] + ("…" if len(data.recommendation) > 800 else "")
    if not text:
        text = "No work-order proposal was created for this investigation."
    _textbox(slide, 0.7, 1.5, 8.6, 3.2, text, size=BODY_FONT_SIZE, color=SLATE)
    if data.recommendation:
        meta = f"Action: {data.recommendation_action}  ·  Risk: {data.recommendation_risk or '—'}"
        _textbox(slide, 0.7, 5.0, 8.6, 0.6, meta, size=BODY_FONT_SIZE, bold=True, color=SLATE)
    _textbox(slide, 0.7, 5.8, 8.6, 0.5,
             "Consequential actions require supervisor approval before execution.",
             size=SMALL_FONT_SIZE, color=GRAY)


def _slide_appendix(prs: PresentationType, data: ReportData) -> None:
    slide = _new_slide(prs)
    _textbox(slide, 0.7, 0.5, 8.6, 0.7, "Evidence Appendix",
             size=TITLE_FONT_SIZE, bold=True, color=SLATE)
    lines = []
    for idx, citation in enumerate(data.citations[:12], start=1):
        title = str(citation.get("title", ""))
        revision = str(citation.get("revision", ""))
        locator = str(citation.get("locator", ""))
        chunk = str(citation.get("chunk_id", ""))
        lines.append(f"[{idx}] {title} · Rev {revision} · {locator} · {chunk}")
    if data.assumptions:
        lines.append("")
        lines.append("Assumptions (labeled, not evidenced):")
        for assumption in data.assumptions[:5]:
            lines.append(f"• {assumption[:140]}")
    if not lines:
        lines.append("No citations recorded.")
    _textbox(slide, 0.7, 1.5, 8.6, 4.8, "\n".join(lines), size=Pt(10), color=SLATE)


def build_pptx(data: ReportData, out_path: str) -> str:
    prs = Presentation()
    _slide_cover(prs, data)
    _slide_finding(prs, data)
    _slide_supporting(prs, data)
    _slide_contradiction(prs, data)
    _slide_recommended(prs, data)
    _slide_appendix(prs, data)
    prs.save(out_path)
    return out_path
