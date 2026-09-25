"""Contract tests for the export endpoints (E1/E2)."""

from __future__ import annotations

import json
from typing import Any

from fastapi.testclient import TestClient


def _create_completed_job(client: TestClient, auth_headers: Any) -> str:
    headers = auth_headers("engineer")
    response = client.post(
        "/api/jobs",
        json={"objective": "Investigate why Pump P-204 failed repeatedly.", "equipment_id": "P-204"},
        headers=headers,
    )
    assert response.status_code == 201
    job_id = str(response.json()["job_id"])
    # drain the SSE stream — the fake graph backend completes the run
    with client.stream("GET", f"/api/jobs/{job_id}/events", headers=headers) as stream:
        for line in stream.iter_lines():
            if not line.startswith("data: "):
                continue
            if json.loads(line[len("data: ") :])["type"] == "completed":
                break
    detail = client.get(f"/api/jobs/{job_id}", headers=headers).json()
    assert detail["status"] == "completed"
    return job_id


def test_export_pptx_builds_deck(client: TestClient, auth_headers: Any, tmp_path: Any) -> None:
    from pptx import Presentation

    job_id = _create_completed_job(client, auth_headers)
    headers = auth_headers("engineer")
    response = client.get(f"/api/jobs/{job_id}/export/pptx", headers=headers)
    assert response.status_code == 200
    assert response.content[:2] == b"PK"  # zip container = valid pptx
    assert "presentationml" in response.headers["content-type"]

    deck_path = tmp_path / "deck.pptx"
    deck_path.write_bytes(response.content)
    deck = Presentation(str(deck_path))
    assert len(deck.slides) == 6
    cover_text = "\n".join(
        shape.text for shape in deck.slides[0].shapes if shape.has_text_frame
    )
    assert "Investigation Report" in cover_text


def test_export_pptx_requires_owner_or_staff(client: TestClient, auth_headers: Any) -> None:
    from app.db import SessionLocal
    from app.models import User
    from app.security import hash_password

    # second (non-owner) engineer, seeded directly like the jobs boundary test
    db = SessionLocal()
    try:
        if db.query(User).filter(User.username == "engineer2").first() is None:
            db.add(
                User(
                    id="00000000-0000-0000-0000-00000000e002",
                    username="engineer2",
                    password_hash=hash_password("nirva-demo-pass"),
                    display_name="Second Engineer",
                    role="engineer",
                )
            )
            db.commit()
    finally:
        db.close()

    job_id = _create_completed_job(client, auth_headers)
    eng2 = {
        "Authorization": "Bearer "
        + client.post(
            "/api/auth/login",
            json={"username": "engineer2", "password": "nirva-demo-pass"},
        ).json()["access_token"]
    }
    assert client.get(f"/api/jobs/{job_id}/export/pptx", headers=eng2).status_code == 403
    # unauthenticated cannot export
    anon = client.get(f"/api/jobs/{job_id}/export/pptx")
    assert anon.status_code == 401
    # supervisor may export any investigation
    sup = client.get(f"/api/jobs/{job_id}/export/pptx", headers=auth_headers("supervisor"))
    assert sup.status_code == 200


def test_export_unknown_job_404(client: TestClient, auth_headers: Any) -> None:
    response = client.get("/api/jobs/does-not-exist/export/pptx", headers=auth_headers("admin"))
    assert response.status_code == 404


def _export_ok(client: TestClient, auth_headers: Any, job_id: str, fmt: str, magic: bytes, ctype: str) -> bytes:
    response = client.get(f"/api/jobs/{job_id}/export/{fmt}", headers=auth_headers("engineer"))
    assert response.status_code == 200, response.text[:200]
    assert response.content[:2] == magic
    assert ctype in response.headers["content-type"]
    return response.content


def test_export_word_and_excel_build(client: TestClient, auth_headers: Any, tmp_path: Any) -> None:
    from docx import Document as DocxDocument
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    job_id = _create_completed_job(client, auth_headers)

    word_bytes = _export_ok(
        client, auth_headers, job_id, "word", b"PK",
        "wordprocessingml.document",
    )
    word_path = tmp_path / "report.docx"
    word_path.write_bytes(word_bytes)
    doc = DocxDocument(str(word_path))
    headings = [p.text for p in doc.paragraphs if (p.style is not None and p.style.name.startswith("Heading"))]
    assert "Primary Finding" in headings
    assert "Evidence Appendix" in headings

    excel_bytes = _export_ok(
        client, auth_headers, job_id, "excel", b"PK",
        "spreadsheetml.sheet",
    )
    excel_path = tmp_path / "report.xlsx"
    excel_path.write_bytes(excel_bytes)
    wb = load_workbook(str(excel_path))
    assert wb.sheetnames == [
        "1_Summary", "2_Evidence", "3_Citations",
        "4_Contradictions", "7_Asset History", "9_WorkingCode",
    ]
    summary = wb["1_Summary"]
    formulas = [str(c.value) for row in summary.iter_rows() for c in row if str(c.value).startswith("=")]
    assert any("COUNTA" in f for f in formulas)
    evidence = wb["2_Evidence"]
    assert evidence.cell(row=1, column=1).value == "No"
    # the fake graph backend stores no evidence rows; live runs fill them
    # (verified against the dev server during E2 development)


def test_documents_export_excel(client: TestClient, auth_headers: Any, tmp_path: Any) -> None:
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    response = client.get("/api/documents/export/excel", headers=auth_headers("engineer"))
    assert response.status_code == 200, response.text[:200]
    assert response.content[:2] == b"PK"
    path = tmp_path / "docs.xlsx"
    path.write_bytes(response.content)
    wb = load_workbook(str(path))
    assert wb.sheetnames == ["Documents"]
    assert wb["Documents"].cell(row=1, column=1).value == "Title"
